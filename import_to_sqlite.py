"""
Import sedan_sales_data.xlsx into sedan_data.db (SQLite).
Run from ~/dev/sedan-exodus/ with sedan_sales_data.xlsx in the same folder.
"""
import sqlite3
import openpyxl

XLSX_PATH = "sedan_sales_data.xlsx"
DB_PATH = "sedan_data.db"

wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws = wb["Sedan Sales Data"]

rows = list(ws.iter_rows(min_row=2, values_only=True))
headers = [c.value for c in ws[1]]  # year, brand, model, body_style, tier, units_sold_us, yoy_us_pct, units_sold_canada, yoy_canada_pct

conn = sqlite3.connect(DB_PATH)
cur = conn.cursor()

cur.execute("DROP TABLE IF EXISTS sedan_sales")
cur.execute("""
CREATE TABLE sedan_sales (
    year INTEGER,
    brand TEXT,
    model TEXT,
    body_style TEXT,
    tier TEXT,
    units_sold_us INTEGER,
    yoy_us_pct REAL,
    units_sold_canada INTEGER,
    yoy_canada_pct REAL
)
""")

def clean(val):
    return None if val == "N/A" else val

insert_rows = []
for r in rows:
    insert_rows.append(tuple(clean(v) for v in r))

cur.executemany(
    "INSERT INTO sedan_sales (year, brand, model, body_style, tier, units_sold_us, yoy_us_pct, units_sold_canada, yoy_canada_pct) VALUES (?,?,?,?,?,?,?,?,?)",
    insert_rows
)

conn.commit()

# quick sanity check printed to console
cur.execute("SELECT COUNT(*), COUNT(DISTINCT model), MIN(year), MAX(year) FROM sedan_sales")
count, models, min_y, max_y = cur.fetchone()
print(f"Imported {count} rows, {models} distinct models, years {min_y}-{max_y}")

cur.execute("SELECT brand, model, COUNT(*) as rows, MIN(year), MAX(year) FROM sedan_sales GROUP BY brand, model ORDER BY brand, model")
print("\nPer-model row counts:")
for row in cur.fetchall():
    print(f"  {row[0]:15} {row[1]:12} {row[2]:3} rows  {row[3]}-{row[4]}")

conn.close()
print(f"\nDone. Database written to {DB_PATH}")
